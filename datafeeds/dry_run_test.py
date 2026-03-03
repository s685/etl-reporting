# ********************************************************************************
# DRY RUN TEST for yaml_multi_worksheet_template.py
# Tests FileWriter and YAML config WITHOUT a Snowflake connection.
# Uses synthetic sample data matching the Application Activity Report structure.
#
# RUN:
#   cd datafeeds/
#   python dry_run_test.py
#
# SUCCESS: Prints "DRY RUN PASSED" and saves output to /tmp/dry_run_output.xlsx
# ********************************************************************************

import os
import sys
import logging
import re
import yaml
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, Color
from datetime import datetime
from typing import cast

logging.basicConfig(
    format='%(asctime)s - %(levelname)s - %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S',
    level=logging.INFO
)

# ---- Inline copies of FileWriter and sanitize_sheet_name (no Snowflake deps) ----
# These are identical to the classes in yaml_multi_worksheet_template.py

def sanitize_sheet_name(name):
    if not name or str(name).strip() == '':
        return 'Sheet'
    name = str(name).strip()
    name = re.sub(r'[\\/*?\[\]:]', '', name)
    if len(name) > 31:
        name = name[:31]
    return name if name else 'Sheet'


class FileWriter:

    def __init__(self, params):
        self.output_path = params["output_path"]
        self.output_file = params["output_file"]
        self.max_column_width = params["max_column_width"]
        self.sheet_header_font = params["sheet_header_font"]
        self.table_header_font = params["table_header_font"]
        self.table_data_font = params["table_data_font"]
        self.group_name_font = params.get("group_name_font")
        self.border_to_row = params["border_to_row"]
        self.carrier_name = params["carrier_name"]
        self.report_name = params["report_name"]
        self.report_start_dt = params["report_start_dt"]
        self.report_end_dt = params["report_end_dt"]
        self.report_run_dt = params["report_run_dt"]
        self.report_as_of_run_dt = params["report_as_of_run_dt"]
        self.header = params["header"]
        self.footer = params["footer"]
        self.group_name_row = params.get("group_name_row")
        self.multi_level_headers = params.get("multi_level_headers")
        self.percent_columns = params.get("percent_columns")
        self.dollar_columns = params["dollar_columns"]
        self.specific_column_widths = params["specific_column_widths"]
        self.positive_dollar_format = "${:,.2f}"
        self.negative_dollar_format = "(${:,.2f})"

    def write_to_excel(self, data, ws, current_page, total_pages, group_name=None):
        last_column = data.shape[1]
        current_row = 1

        if self.header:
            self.add_header(ws, current_row, last_column, current_page, total_pages)
            current_row += 4

        if self.group_name_row and group_name is not None:
            self.write_group_name(ws, current_row, last_column, group_name)
            current_row += 1

        if self.multi_level_headers:
            current_row = self.write_multi_level_headers(ws, current_row, last_column)
        else:
            name, size, bold, color, alignment, wrap_text, fill_color, fill_type = self.set_cell_properties(self.table_header_font)
            for col, header in enumerate(data.columns, start=1):
                cell = ws.cell(row=current_row, column=col)
                cell.value = header
                cell.font = Font(name=name, size=size, bold=bold, color=color)
                cell.alignment = Alignment(horizontal=alignment, wrap_text=wrap_text)
                cell.fill = PatternFill(fill_type=fill_type, fgColor=fill_color)
            current_row += 1

        logging.info("Setting column widths")
        self.set_column_widths(ws, self.max_column_width)
        self.set_specific_column_widths(data, ws)

        data_start_row = current_row
        data_rows = data.values.tolist()

        name, size, bold, color, alignment, wrap_text, fill_color, fill_type = self.set_cell_properties(self.table_data_font)

        for row in data_rows:
            for col_idx, val in enumerate(row, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = Font(name=name, size=size, bold=bold, color=color)
                if col_idx == 1:
                    cell.alignment = Alignment(horizontal='left', wrap_text=wrap_text)
                else:
                    cell.alignment = Alignment(horizontal=alignment, wrap_text=wrap_text)
                if fill_type and fill_type != 'none':
                    cell.fill = PatternFill(fill_type=fill_type, fgColor=fill_color)
            current_row += 1

        self.apply_percent_format(data, ws, data_start_row)

    def write_group_name(self, ws, current_row, last_column, group_name):
        name, size, bold, color, alignment, wrap_text, fill_color, fill_type = self.set_cell_properties(
            self.group_name_font or self.table_header_font
        )
        cell = ws.cell(row=current_row, column=1, value=str(group_name))
        cell.font = Font(name=name, size=size, bold=bold, color=color)
        cell.alignment = Alignment(horizontal=alignment, wrap_text=wrap_text)
        cell.fill = PatternFill(fill_type=fill_type, fgColor=fill_color)
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=last_column)
        for col in range(2, last_column + 1):
            ws.cell(row=current_row, column=col).fill = PatternFill(fill_type=fill_type, fgColor=fill_color)

    def write_multi_level_headers(self, ws, start_row, last_column):
        name, size, bold, color, alignment, wrap_text, fill_color, fill_type = self.set_cell_properties(self.table_header_font)
        thin_border = Border(
            top=Side(style='thin', color='000000'),
            bottom=Side(style='thin', color='000000'),
            left=Side(style='thin', color='000000'),
            right=Side(style='thin', color='000000'),
        )

        current_row = start_row
        for level in self.multi_level_headers:
            col_offset = 1
            for cell_def in level:
                label = cell_def.get('label', '')
                span = cell_def.get('span', 1)

                cell = ws.cell(row=current_row, column=col_offset, value=label)
                cell.font = Font(name=name, size=size, bold=bold, color=color)
                cell.alignment = Alignment(horizontal=alignment, wrap_text=wrap_text, vertical='center')
                cell.fill = PatternFill(fill_type=fill_type, fgColor=fill_color)
                cell.border = thin_border

                if span > 1:
                    ws.merge_cells(
                        start_row=current_row, start_column=col_offset,
                        end_row=current_row, end_column=col_offset + span - 1
                    )
                    for c in range(col_offset + 1, col_offset + span):
                        merged_cell = ws.cell(row=current_row, column=c)
                        merged_cell.fill = PatternFill(fill_type=fill_type, fgColor=fill_color)
                        merged_cell.border = thin_border

                col_offset += span
            current_row += 1

        return current_row

    def apply_percent_format(self, data, ws, data_start_row):
        if not self.percent_columns:
            return
        for col_name in self.percent_columns:
            if col_name in data.columns:
                col_idx = cast(int, data.columns.get_loc(col_name)) + 1
                for row_idx in range(data_start_row, data_start_row + len(data)):
                    ws.cell(row=row_idx, column=col_idx).number_format = '0.0%'
                logging.info(f"Applied percent format to column: {col_name}")

    def apply_dollar_format(self, data, ws):
        if isinstance(data, pd.DataFrame):
            for column in self.dollar_columns:
                logging.info(f"Applying dollar format to column: {column}")
                if column in data.columns:
                    col_idx = cast(int, data.columns.get_loc(column)) + 1
                    column_letter = get_column_letter(col_idx)
                    for cell in ws[column_letter]:
                        cell.number_format = '$#,##0.00'
        else:
            data = self.positive_dollar_format.format(data) if data > 0 else self.negative_dollar_format.format(abs(data))
        return data

    def apply_border(self, ws, current_row, last_column, border_to_row):
        if border_to_row and border_to_row.get('border_to_table_headers'):
            self.apply_border_to_row(ws, current_row, last_column, border_to_row)
        return current_row + 1

    def apply_border_to_row(self, ws, current_row, last_column, border_to_row):
        start_color = border_to_row['start_color']
        end_color = border_to_row['end_color']
        fill_type = border_to_row['fill_type']
        thin_border = Border(top=Side(style='thin'), bottom=Side(style='thin'))
        ws.row_dimensions[current_row].height = 1
        for col_num in range(1, last_column + 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.border = thin_border
            cell.fill = PatternFill(start_color=Color(rgb=start_color), end_color=Color(rgb=end_color), fill_type=fill_type)
        current_row += 1
        return current_row

    def add_header(self, ws, current_row, last_column, current_page, total_pages):
        name, size, bold, color, alignment, wrap_text, fill_color, fill_type = self.set_cell_properties(self.sheet_header_font)

        for row in range(current_row, current_row + 3):
            cell = ws.cell(row=row, column=1)
            if row == current_row:
                cell.value = self.carrier_name
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column // 2)
                cell.font = Font(name=name, size=size, bold=bold, color=color)
                cell.alignment = Alignment(horizontal=alignment, wrap_text=wrap_text)
                cell.fill = PatternFill(fill_type=fill_type, fgColor=fill_color)

                time_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                time_info = f"Executed On: {time_str}"
                cell_offset = last_column // 2
                time_cell = ws.cell(row=row, column=cell_offset + 1)
                time_cell.value = time_info
                ws.merge_cells(start_row=row, start_column=cell_offset + 1, end_row=row, end_column=last_column)
                time_cell.font = Font(name=name, size=size, bold=bold, color=color)
                time_cell.alignment = Alignment(horizontal='right', wrap_text=wrap_text)
                time_cell.fill = PatternFill(fill_type=fill_type, fgColor=fill_color)

            elif row == current_row + 1:
                cell.value = self.report_name
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column // 2)
                cell.font = Font(name=name, size=size, bold=bold, color=color)
                cell.alignment = Alignment(horizontal=alignment, wrap_text=wrap_text)
                cell.fill = PatternFill(fill_type=fill_type, fgColor=fill_color)

                page_info = f"Page {current_page} of {total_pages}"
                cell_offset = last_column // 2
                page_cell = ws.cell(row=row, column=cell_offset + 1)
                page_cell.value = page_info
                ws.merge_cells(start_row=row, start_column=cell_offset + 1, end_row=row, end_column=last_column)
                page_cell.font = Font(name=name, size=size, bold=bold, color=color)
                page_cell.alignment = Alignment(horizontal='right', wrap_text=wrap_text)
                page_cell.fill = PatternFill(fill_type=fill_type, fgColor=fill_color)
            else:
                if self.report_start_dt and self.report_end_dt:
                    try:
                        start_date = datetime.strptime(self.report_start_dt, '%Y-%m-%d %H:%M:%S.%f').strftime("%m/%d/%Y")
                    except ValueError:
                        try:
                            start_date = datetime.strptime(self.report_start_dt, '%Y-%m-%d %H:%M:%S').strftime("%m/%d/%Y")
                        except ValueError:
                            start_date = self.report_start_dt
                    try:
                        end_date = datetime.strptime(self.report_end_dt, '%Y-%m-%d %H:%M:%S.%f').strftime("%m/%d/%Y")
                    except ValueError:
                        try:
                            end_date = datetime.strptime(self.report_end_dt, '%Y-%m-%d %H:%M:%S').strftime("%m/%d/%Y")
                        except ValueError:
                            end_date = self.report_end_dt
                    cell.value = f"For Dates {start_date} - {end_date}"
                elif self.report_run_dt:
                    try:
                        report_date = datetime.strptime(self.report_run_dt, '%Y-%m-%d %H:%M:%S.%f').strftime("%m/%d/%Y")
                    except ValueError:
                        try:
                            report_date = datetime.strptime(self.report_run_dt, '%Y-%m-%d %H:%M:%S').strftime("%m/%d/%Y")
                        except ValueError:
                            report_date = self.report_run_dt
                    cell.value = f"Report as Date: {report_date}"
                else:
                    cell.value = f"Report as Date: {datetime.now().strftime('%m/%d/%Y')}"
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
                cell.font = Font(name=name, size=size, bold=bold, color=color)
                cell.alignment = Alignment(horizontal=alignment, wrap_text=wrap_text)
                cell.fill = PatternFill(fill_type=fill_type, fgColor=fill_color)

    def set_column_widths(self, ws, max_column_width):
        width = self.max_column_width or 15
        for col in range(1, ws.max_column + 1):
            col_index = get_column_letter(col)
            ws.column_dimensions[col_index].width = width

    def set_specific_column_widths(self, data, ws):
        if isinstance(data, pd.DataFrame):
            if self.specific_column_widths is not None:
                logging.info("Setting specific column widths")
                for column in self.specific_column_widths:
                    clmn = column['column']
                    wdth = column['width']
                    ws.column_dimensions[clmn].width = wdth

    def set_cell_properties(self, font):
        if not font or not isinstance(font, dict):
            return 'Calibri', 11, False, '000000', 'left', False, 'FFFFFF', 'solid'
        name = font['name']
        size = font['size']
        bold = font['bold']
        color = font['color']
        wrap_text = font['wrap_text']
        fill_color = font['fill_color']
        fill_type = font['fill_type']
        alignment = font['alignment']
        return name, size, bold, color, alignment, wrap_text, fill_color, fill_type


# ---- DRY RUN ----

def run_dry_run():
    print("\n" + "=" * 70)
    print("  DRY RUN: yaml_multi_worksheet_template.py")
    print("=" * 70)

    config_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'application_activity_report.yml')

    # 1. Load and validate YAML
    print("\n[1] Loading YAML config...")
    with open(config_path, 'r') as f:
        config = yaml.safe_load(f)

    assert config.get('carrier_name'), "FAIL: carrier_name missing"
    assert config.get('report_name'), "FAIL: report_name missing"
    assert config.get('table'), "FAIL: table missing"
    assert config.get('grouping_column'), "FAIL: grouping_column missing"
    assert config.get('pre_sql_query'), "FAIL: pre_sql_query missing"
    assert isinstance(config.get('multi_level_headers'), list), "FAIL: multi_level_headers must be a list"
    print(f"    carrier_name    : {config['carrier_name']}")
    print(f"    report_name     : {config['report_name']}")
    print(f"    table           : {config['table']}")
    print(f"    grouping_column : {config['grouping_column']}")
    print(f"    header levels   : {len(config['multi_level_headers'])}")
    print("    YAML: OK")

    # 2. Validate multi-level header spans
    print("\n[2] Validating multi_level_headers span totals...")
    EXPECTED_COLUMNS = 26  # 1 label + 6+6+6+6 age/gender + 1 Total
    for i, level in enumerate(config['multi_level_headers']):
        total_span = sum(cell_def['span'] for cell_def in level)
        assert total_span == EXPECTED_COLUMNS, (
            f"FAIL: Level {i+1} span total={total_span}, expected={EXPECTED_COLUMNS}"
        )
        print(f"    Level {i+1}: {total_span} columns  OK")

    # 3. Build synthetic sample data (26 columns after removing GROUP_NAME)
    print("\n[3] Creating synthetic sample data...")
    grouping_column = config['grouping_column']

    # Column names matching the YAML percent_columns and the report structure
    data_columns = [
        grouping_column,
        'APPLICATION_DESCRIPTION',
        'AGE_LT65_MALE_NO',    'AGE_LT65_MALE_PCT',
        'AGE_LT65_FEMALE_NO',  'AGE_LT65_FEMALE_PCT',
        'AGE_LT65_ND_NO',      'AGE_LT65_ND_PCT',
        'AGE_65_74_MALE_NO',   'AGE_65_74_MALE_PCT',
        'AGE_65_74_FEMALE_NO', 'AGE_65_74_FEMALE_PCT',
        'AGE_65_74_ND_NO',     'AGE_65_74_ND_PCT',
        'AGE_75PLUS_MALE_NO',  'AGE_75PLUS_MALE_PCT',
        'AGE_75PLUS_FEMALE_NO','AGE_75PLUS_FEMALE_PCT',
        'AGE_75PLUS_ND_NO',    'AGE_75PLUS_ND_PCT',
        'ALL_MALE_NO',         'ALL_MALE_PCT',
        'ALL_FEMALE_NO',       'ALL_FEMALE_PCT',
        'ALL_ND_NO',           'ALL_ND_PCT',
        'TOTAL',
    ]

    row_labels = [
        'Applications In Process at Period Start',
        'New Applications',
        'Reopened Applications',
        'Applications Completed',
        '  Approved',
        '  Declined',
        '    Straight Decline',
        '    Auto Decline',
        '  Withdrawn',
        '  Withdrawn - Entered In Error',
        'Incomplete',
        'Applications In Process at Period End',
    ]

    # Build rows: 2 groups × 12 labels
    groups = ['NYL My Care', 'New York Life Insurance Company Totals']
    rows = []
    for g_idx, group in enumerate(groups):
        for r_idx, label in enumerate(row_labels):
            # Realistic sample values: some non-zero counts and percentages
            base = (g_idx + 1) * (r_idx + 1)
            row = [group, label] + [base, round(base * 0.01, 4)] * 12 + [base * 3]
            rows.append(row)

    df = pd.DataFrame(rows, columns=data_columns)
    print(f"    Rows: {len(df)}  |  Columns: {len(df.columns)}  (including {grouping_column})")
    print(f"    Groups: {df[grouping_column].unique().tolist()}")

    # Verify percent_columns exist in sample data
    percent_cols = config.get('percent_columns', [])
    for pc in percent_cols:
        assert pc in df.columns, f"FAIL: percent_column '{pc}' not in sample data columns"
    print(f"    All {len(percent_cols)} percent_columns found in sample data  OK")

    # 4. Build params dict (same as main())
    print("\n[4] Building FileWriter params...")
    output_dir = '/tmp'
    output_file = 'dry_run_output.xlsx'

    params = {
        'output_path': output_dir,
        'output_file': output_file,
        'max_column_width': config.get('max_column_width'),
        'sheet_header_font': config.get('sheet_header_font'),
        'table_header_font': config.get('table_header_font'),
        'table_data_font': config.get('table_data_font'),
        'group_name_font': config.get('group_name_font'),
        'carrier_name': config['carrier_name'],
        'report_name': config['report_name'],
        'report_start_dt': '2026-02-01 00:00:00',
        'report_end_dt': '2026-02-28 00:00:00',
        'report_as_of_run_dt': None,
        'report_run_dt': None,
        'header': config.get('header'),
        'footer': config.get('footer'),
        'group_name_row': config.get('group_name_row'),
        'multi_level_headers': config.get('multi_level_headers'),
        'percent_columns': percent_cols,
        'border_to_row': config.get('border_to_row'),
        'dollar_columns': config.get('dollar_columns') or [],
        'specific_column_widths': config.get('specific_column_widths'),
    }
    print("    Params built  OK")

    # 5. Run FileWriter for each group (mirrors main() loop)
    print("\n[5] Running FileWriter for each group...")
    null_group_sheet_name = config.get('null_group_sheet_name', 'Unassigned')
    null_count = df[grouping_column].isna().sum()
    if null_count > 0:
        df[grouping_column] = df[grouping_column].fillna(null_group_sheet_name)

    unique_groups = df[grouping_column].unique().tolist()
    total_pages = len(unique_groups)

    wb = Workbook()
    del wb['Sheet']
    writer = FileWriter(params)

    for current_page, group_value in enumerate(unique_groups, start=1):
        sheet_name = sanitize_sheet_name(group_value)
        print(f"    [{current_page}/{total_pages}] Writing sheet: '{sheet_name}'")

        group_df = df[df[grouping_column] == group_value].drop(columns=[grouping_column])
        group_df = group_df.reset_index(drop=True)

        ws = wb.create_sheet(title=sheet_name)
        writer.write_to_excel(group_df, ws, current_page, total_pages, group_value)

        dollar_columns = config.get('dollar_columns')
        if dollar_columns:
            writer.apply_dollar_format(group_df, ws)

    # 6. Save workbook
    output_path = os.path.join(output_dir, output_file)
    wb.save(output_path)
    print(f"\n[6] File saved: {output_path}")

    # 7. Verify output
    print("\n[7] Verifying output...")
    assert os.path.exists(output_path), f"FAIL: Output file not found at {output_path}"

    wb_check = load_workbook(output_path)
    sheet_names = wb_check.sheetnames
    print(f"    Worksheets    : {sheet_names}")
    assert len(sheet_names) == len(unique_groups), \
        f"FAIL: expected {len(unique_groups)} sheets, got {len(sheet_names)}"

    for expected_group, actual_sheet in zip(unique_groups, sheet_names):
        expected_name = sanitize_sheet_name(expected_group)
        assert actual_sheet == expected_name, \
            f"FAIL: expected sheet '{expected_name}', got '{actual_sheet}'"
        print(f"    Sheet '{actual_sheet}'  OK")

    # Spot-check first sheet structure
    ws0 = wb_check[sheet_names[0]]
    print(f"    Sheet[0] dimensions: {ws0.max_row} rows x {ws0.max_column} cols")

    # With header(3 rows) + spacer(1) + group name(1) + 3 header levels + 12 data rows = 20 rows
    expected_rows = 3 + 1 + 1 + 3 + len(row_labels)  # = 20
    assert ws0.max_row == expected_rows, \
        f"FAIL: expected {expected_rows} rows, got {ws0.max_row}"
    print(f"    Row count = {ws0.max_row} (expected {expected_rows})  OK")

    assert ws0.max_column == EXPECTED_COLUMNS, \
        f"FAIL: expected {EXPECTED_COLUMNS} columns, got {ws0.max_column}"
    print(f"    Column count = {ws0.max_column} (expected {EXPECTED_COLUMNS})  OK")

    # Verify carrier name in row 1
    row1_val = ws0.cell(row=1, column=1).value
    assert row1_val == config['carrier_name'], \
        f"FAIL: row 1 col 1 = '{row1_val}', expected '{config['carrier_name']}'"
    print(f"    Header row 1: '{row1_val}'  OK")

    # Verify report name in row 2
    row2_val = ws0.cell(row=2, column=1).value
    assert row2_val == config['report_name'], \
        f"FAIL: row 2 col 1 = '{row2_val}', expected '{config['report_name']}'"
    print(f"    Header row 2: '{row2_val}'  OK")

    # Verify date row
    row3_val = ws0.cell(row=3, column=1).value
    assert row3_val and 'For Dates' in row3_val, \
        f"FAIL: row 3 col 1 = '{row3_val}', expected 'For Dates ...'"
    print(f"    Header row 3: '{row3_val}'  OK")

    # Verify group name row (row 5)
    row5_val = ws0.cell(row=5, column=1).value
    assert row5_val == unique_groups[0], \
        f"FAIL: group name row = '{row5_val}', expected '{unique_groups[0]}'"
    print(f"    Group name row: '{row5_val}'  OK")

    # Verify level 1 header (row 6): Age < 65 should be at column 2
    row6_col2 = ws0.cell(row=6, column=2).value
    assert row6_col2 == 'Age < 65', \
        f"FAIL: level 1 header col 2 = '{row6_col2}', expected 'Age < 65'"
    print(f"    Level 1 header (Age < 65): OK")

    # Verify level 2 header (row 7): Male at column 2
    row7_col2 = ws0.cell(row=7, column=2).value
    assert row7_col2 == 'Male', \
        f"FAIL: level 2 header col 2 = '{row7_col2}', expected 'Male'"
    print(f"    Level 2 header (Male): OK")

    # Verify level 3 header (row 8): No. at column 2
    row8_col2 = ws0.cell(row=8, column=2).value
    assert row8_col2 == 'No.', \
        f"FAIL: level 3 header col 2 = '{row8_col2}', expected 'No.'"
    print(f"    Level 3 header (No.): OK")

    # Verify page info in row 2 right side
    page_cell = ws0.cell(row=2, column=ws0.max_column // 2 + 1).value
    assert page_cell and 'Page' in str(page_cell), \
        f"FAIL: page info = '{page_cell}'"
    print(f"    Page info: '{page_cell}'  OK")

    print("\n" + "=" * 70)
    print("  DRY RUN PASSED - No errors! File ready for use.")
    print(f"  Output: {output_path}")
    print("=" * 70 + "\n")


if __name__ == '__main__':
    run_dry_run()
